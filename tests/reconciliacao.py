#!/usr/bin/env python3
"""
Teste 12 — reconciliação com a planilha.

Lança os MESMOS dados de exemplo (a) numa cópia da planilha original e (b) no aplicativo,
recalcula a planilha com o LibreOffice e compara os totais.

Uso:  python3 tests/reconciliacao.py "<planilha original.xlsx>"
Requer: LibreOffice (soffice) e o script de recálculo em /mnt/skills/public/xlsx/scripts/recalc.py
        (ou RECALC=/caminho/recalc.py).
"""
import datetime as dt, json, os, shutil, subprocess, sys, tempfile, time, urllib.request
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT.parent / "analise" / "fonte.xlsx")
RECALC = os.environ.get("RECALC", "/mnt/skills/public/xlsx/scripts/recalc.py")
PORT = 3198

# ---------------- dados de exemplo (agosto/2026) ----------------
ATEND = [  # data, barbeiro, serviço, forma, valor recebido
    (dt.date(2026, 8, 3), "Pedro", "Corte", "Pix", 55),
    (dt.date(2026, 8, 3), "André", "Corte e Barba", "Cartão de Crédito", 96.5),
    (dt.date(2026, 8, 4), "Christopher", "Manutenção Prótese Capilar", "Cartão de Débito", 107.8),
    (dt.date(2026, 8, 5), "Pedro", "Barba", "Dinheiro", 45),
    (dt.date(2026, 8, 6), "André", "Corte e Barboterapia", "Pix", 110),
]
BALCAO = [  # data, barbeiro, produto(nome), sku, qtd, preço unit., forma
    (dt.date(2026, 8, 3), "Pedro", "Óleo Argan", "ARG-001", 2, 60, "Pix"),
    (dt.date(2026, 8, 5), "André", "Fixa Gold", "FTG-001", 1, 45, "Dinheiro"),
    (dt.date(2026, 8, 6), "Christopher", "Shampoo Doux", "SHD-001", 3, 75, "Cartão de Crédito"),
]
ML = [  # data, sku, qtd, preço, plano, %comissão, frete
    (dt.date(2026, 8, 4), "PC-050", 1, 581.22, "Premium", 0.15, 22.55),
    (dt.date(2026, 8, 7), "FEB-001", 3, 16.14, "Clássico", 0.17, 0),
]
SHOPEE = [  # data, sku, qtd, preço, tipo, alto volume, %comissão, frete
    (dt.date(2026, 8, 8), "MAS-001", 2, 85, "CNPJ", "Não", 0.14, 0),
]
SAIDAS = [  # data, descrição, categoria, valor, forma
    (dt.date(2026, 8, 3), "Café e copos", "Suprimentos", 42.9, "Dinheiro"),
    (dt.date(2026, 8, 6), "Lâminas", "Suprimentos", 58, "Pix"),
]


def planilha():
    tmp = Path(tempfile.mkdtemp())
    f = tmp / "reconc.xlsx"
    shutil.copy(SRC, f)
    wb = openpyxl.load_workbook(f)
    ag = wb["Agosto 2026"]
    for i, (d, b, s, fp, rec) in enumerate(ATEND):
        r = 7 + i
        ag.cell(r, 1, d); ag.cell(r, 2, b); ag.cell(r, 3, s); ag.cell(r, 5, fp); ag.cell(r, 6, rec)
    for i, (d, desc, cat, v, fp) in enumerate(SAIDAS):
        r = 7 + i
        ag.cell(r, 11, d); ag.cell(r, 12, desc); ag.cell(r, 13, cat); ag.cell(r, 14, v); ag.cell(r, 15, fp)
    bc = wb["Barbeiro e Comissão"]
    for i, (d, b, nome, _sku, q, p, fp) in enumerate(BALCAO):
        r = 12 + i
        bc.cell(r, 1, d); bc.cell(r, 2, b); bc.cell(r, 3, nome); bc.cell(r, 4, q); bc.cell(r, 5, p); bc.cell(r, 6, fp)
    ml = wb["Vendas Mercado Livre"]
    for i, (d, sku, q, p, plano, pct, frete) in enumerate(ML):
        r = 2 + i
        ml.cell(r, 1, d); ml.cell(r, 2, sku); ml.cell(r, 4, q); ml.cell(r, 5, p); ml.cell(r, 6, plano); ml.cell(r, 8, pct); ml.cell(r, 11, frete); ml.cell(r, 15, "Entregue")
    sh = wb["Vendas Shopee"]
    for i, (d, sku, q, p, tipo, alto, pct, frete) in enumerate(SHOPEE):
        r = 2 + i
        sh.cell(r, 1, d); sh.cell(r, 2, sku); sh.cell(r, 4, q); sh.cell(r, 5, p); sh.cell(r, 6, tipo); sh.cell(r, 7, alto); sh.cell(r, 9, pct); sh.cell(r, 13, frete); sh.cell(r, 17, "Entregue")
    wb.save(f)
    out = subprocess.run([sys.executable, RECALC, str(f), "240"], capture_output=True, text=True)
    info = json.loads(out.stdout[out.stdout.index("{"):])
    assert info.get("total_errors", 0) == 0, info
    v = openpyxl.load_workbook(f, data_only=True)
    rm, bcv, mlv, shv, pv = v["Resumo Mensal"], v["Barbeiro e Comissão"], v["Vendas Mercado Livre"], v["Vendas Shopee"], v["Produtos"]
    cents = lambda x: round(float(x or 0) * 100)
    ml_rows = [r for r in range(2, 2 + len(ML))]
    sh_rows = [r for r in range(2, 2 + len(SHOPEE))]
    res = {
        "servicos_valor_cheio": cents(rm["C6"].value), "servicos_recebido": cents(rm["D6"].value),
        "comissao": cents(rm["F6"].value), "saidas": cents(rm["I6"].value),
        "balcao_total": cents(bcv["C5"].value), "balcao_qtd": int(bcv["C9"].value or 0),
        "ml_bruto": sum(cents(mlv.cell(r, 12).value) for r in ml_rows),
        "ml_taxas": sum(cents(mlv.cell(r, 9).value) + cents(mlv.cell(r, 10).value) for r in ml_rows),
        "ml_frete": sum(cents(mlv.cell(r, 11).value) for r in ml_rows),
        "ml_liquido": sum(cents(mlv.cell(r, 13).value) for r in ml_rows),
        "shopee_bruto": sum(cents(shv.cell(r, 14).value) for r in sh_rows),
        "shopee_liquido": sum(cents(shv.cell(r, 15).value) for r in sh_rows),
        "_ml_linhas": [(cents(mlv.cell(r, 12).value), cents(mlv.cell(r, 9).value) + cents(mlv.cell(r, 10).value), cents(mlv.cell(r, 11).value)) for r in ml_rows],
        "_sh_linhas": [(cents(shv.cell(r, 14).value), cents(shv.cell(r, 11).value) + cents(shv.cell(r, 12).value), cents(shv.cell(r, 13).value)) for r in sh_rows],
        "estoque": {},
    }
    for r in range(2, 60):
        sku = pv.cell(r, 1).value
        if sku:
            res["estoque"][sku] = int((pv.cell(r, 6).value or 0) - (pv.cell(r, 13).value or 0))
    shutil.rmtree(tmp, ignore_errors=True)
    return res


class App:
    def __init__(self):
        self.tmp = Path(tempfile.mkdtemp())
        env = {**os.environ, "DB_FILE": str(self.tmp / "r.db"), "PORT": str(PORT), "VERDUGO_EMAIL": "dono@verdugo.test", "VERDUGO_SENHA": "SenhaForte2026"}
        subprocess.run(["npm", "run", "-s", "definir-senha"], cwd=ROOT, env=env, check=True, capture_output=True)
        subprocess.run(["npm", "run", "-s", "importar-planilha", "--", "dados/planilha_extraida.json", env["DB_FILE"]], cwd=ROOT, env=env, check=True, capture_output=True)
        self.proc = subprocess.Popen(["node", "--disable-warning=ExperimentalWarning", "server/index.js"], cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(50):
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{PORT}/saude"); break
            except Exception:
                time.sleep(0.2)
        r = self.req("POST", "/api/login", {"email": "dono@verdugo.test", "senha": "SenhaForte2026"}, raw=True)
        self.cookie = r.headers["Set-Cookie"].split(";")[0]

    def req(self, method, path, body=None, raw=False):
        data = json.dumps(body).encode() if body is not None else None
        h = {"x-verdugo": "1", "content-type": "application/json"}
        if getattr(self, "cookie", None):
            h["cookie"] = self.cookie
        rq = urllib.request.Request(f"http://127.0.0.1:{PORT}{path}", data=data, method=method, headers=h)
        resp = urllib.request.urlopen(rq)
        return resp if raw else json.loads(resp.read())

    def close(self):
        self.proc.terminate(); self.proc.wait(); shutil.rmtree(self.tmp, ignore_errors=True)


def lancar_no_app(app, plan):
    barbers = {b["name"]: b["id"] for b in app.req("GET", "/api/barbeiros")}
    servs = {s["name"]: s["id"] for s in app.req("GET", "/api/servicos")}
    for d, b, s, fp, rec in ATEND:
        app.req("POST", "/api/atendimentos", {"date": d.isoformat(), "barber_id": barbers[b], "payment_method": fp,
                                              "items": [{"service_id": servs[s], "is_main": True}], "received": str(rec)})
    for d, desc, cat, v, fp in SAIDAS:
        app.req("POST", "/api/saidas", {"date": d.isoformat(), "description": desc, "category": cat, "amount": str(v), "payment_method": fp})
    for d, b, _nome, sku, q, p, fp in BALCAO:
        app.req("POST", "/api/balcao", {"date": d.isoformat(), "seller_barber_id": barbers[b], "payment_method": fp,
                                        "items": [{"sku": sku, "qty": q, "price": str(p)}]})
    # nos pedidos online o app recebe as taxas EFETIVAMENTE cobradas (aqui: as mesmas que a planilha calculou)
    for i, (d, sku, q, p, *_rest) in enumerate(ML):
        bruto, taxas, frete = plan["_ml_linhas"][i]
        app.req("POST", "/api/online", {"channel": "mercado_livre", "order_number": f"REC-ML-{i}", "date": d.isoformat(), "status": "entregue",
                                        "items": [{"sku": sku, "qty": q, "unit_price": str(p)}], "gross": bruto / 100, "fees": taxas / 100, "shipping": frete / 100})
    for i, (d, sku, q, p, *_rest) in enumerate(SHOPEE):
        bruto, taxas, frete = plan["_sh_linhas"][i]
        app.req("POST", "/api/online", {"channel": "shopee", "order_number": f"REC-SH-{i}", "date": d.isoformat(), "status": "entregue",
                                        "items": [{"sku": sku, "qty": q, "unit_price": str(p)}], "gross": bruto / 100, "fees": taxas / 100, "shipping": frete / 100})
    rep = app.req("GET", "/api/relatorio?periodo=mes&ref=2026-08-01")
    prods = {p["sku"]: p["saldo"] for p in app.req("GET", "/api/produtos?todos=1")}
    c = {x["canal"]: x for x in rep["online"]["canais"]}
    return {
        "servicos_valor_cheio": rep["servicos"]["bruto"], "servicos_recebido": rep["servicos"]["recebido"],
        "comissao": rep["servicos"]["comissao"], "saidas": rep["caixa"]["saidas"],
        "balcao_total": rep["balcao"]["total"], "balcao_qtd": rep["balcao"]["itens"],
        "ml_bruto": c["mercado_livre"]["vendas"]["bruto"], "ml_taxas": c["mercado_livre"]["vendas"]["taxas"],
        "ml_frete": c["mercado_livre"]["vendas"]["frete"], "ml_liquido": c["mercado_livre"]["vendas"]["liquido"],
        "shopee_bruto": c["shopee"]["vendas"]["bruto"], "shopee_liquido": c["shopee"]["vendas"]["liquido"],
        "estoque": prods,
    }


if __name__ == "__main__":
    plan = planilha()
    app = App()
    try:
        a = lancar_no_app(app, plan)
    finally:
        app.close()
    linhas, falhas = [], 0
    for k in [k for k in plan if not k.startswith("_") and k != "estoque"]:
        igual = plan[k] == a[k]
        falhas += not igual
        linhas.append(f"{'OK   ' if igual else 'DIF  '} {k:22s} planilha={plan[k]/100 if 'qtd' not in k else plan[k]:>10}  app={a[k]/100 if 'qtd' not in k else a[k]:>10}")
    difs = [s for s in plan["estoque"] if plan["estoque"][s] != a["estoque"].get(s)]
    falhas += len(difs)
    linhas.append(f"{'OK   ' if not difs else 'DIF  '} estoque por SKU        {len(plan['estoque'])} SKUs comparados, {len(difs)} diferença(s) {difs[:5]}")
    print("\n".join(linhas))
    (ROOT / "docs" / "evidencias").mkdir(parents=True, exist_ok=True)
    (ROOT / "docs" / "evidencias" / "reconciliacao.txt").write_text("\n".join(linhas) + "\n")
    print(f"\n{'RECONCILIADO' if not falhas else f'{falhas} DIFERENÇA(S)'}")
    sys.exit(1 if falhas else 0)
